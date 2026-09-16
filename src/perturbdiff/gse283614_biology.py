from __future__ import annotations

from dataclasses import dataclass
import math
from pathlib import Path
from typing import Iterable, Sequence

import numpy as np
import pandas as pd
from scipy import stats

from .gse283614 import (
    SAFE_HARBOR_GUIDES,
    guide_family,
    read_table_s1_grna_library,
    table_s1_name_to_guide_identity,
)


EXPECTED_RESPONSE_MODELS = {
    "ESC_cluster1_score_vs_NANOG": "linear",
    "ESC_cluster2_score_vs_OCT4": "hill",
    "DE_differentiation_score_vs_NANOG": "quadratic",
    "DE_differentiation_score_vs_OCT4": "quadratic",
    "DE_off-target_score_vs_NANOG": "hill",
    "DE_off-target_score_vs_OCT4": "hill",
}


@dataclass(frozen=True)
class ResponseSpec:
    name: str
    target: str
    score_column: str
    expected_model: str


RESPONSE_SPECS = (
    ResponseSpec("ESC_cluster1_score_vs_NANOG", "NANOG", "esc_score", "linear"),
    ResponseSpec("ESC_cluster2_score_vs_OCT4", "OCT4", "esc_score", "hill"),
    ResponseSpec("DE_differentiation_score_vs_NANOG", "NANOG", "de_differentiation_score", "quadratic"),
    ResponseSpec("DE_differentiation_score_vs_OCT4", "OCT4", "de_differentiation_score", "quadratic"),
    ResponseSpec("DE_off-target_score_vs_NANOG", "NANOG", "de_off_target_score", "hill"),
    ResponseSpec("DE_off-target_score_vs_OCT4", "OCT4", "de_off_target_score", "hill"),
)


def _sheet_records(path: str | Path, sheet_name: str) -> tuple[list[str], list[dict[str, object]]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise ImportError("Reading the GSE283614 dosage workbook requires openpyxl") from exc
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise KeyError(f"Sheet {sheet_name!r} not found in {path}")
    rows = list(workbook[sheet_name].iter_rows(values_only=True))
    header_idx = next(
        (idx for idx, row in enumerate(rows) if row and str(row[0] or "").strip() in {"gRNA", "Name"}),
        None,
    )
    if header_idx is None:
        raise ValueError(f"Could not find a gRNA/Name header in sheet {sheet_name!r}")
    headers = [str(value or "").strip().replace("\n", " ") for value in rows[header_idx]]
    records: list[dict[str, object]] = []
    for values in rows[header_idx + 1 :]:
        if not any(value is not None for value in values):
            continue
        records.append(
            {
                headers[idx]: values[idx]
                for idx in range(min(len(headers), len(values)))
                if headers[idx]
            }
        )
    return headers, records


def _author_name_mapping(guide_library_workbook: str | Path) -> dict[str, str]:
    library = read_table_s1_grna_library(guide_library_workbook)
    mapping: dict[str, str] = {}
    for _, row in library.iterrows():
        identity, _ = table_s1_name_to_guide_identity(row)
        name = str(row.get("table_s1_name", "")).strip()
        if name and identity:
            mapping[name] = identity
    for idx in range(1, 8):
        mapping[f"NT-Control-{idx}"] = f"NT_Control_{idx}"
        mapping[f"NT-{idx}"] = f"NT_Control_{idx}"
    for idx, identity in enumerate(SAFE_HARBOR_GUIDES, start=1):
        mapping[f"SafeHarbor-{idx}"] = identity
    return mapping


def read_author_dosage_table(
    dosage_workbook: str | Path,
    *,
    guide_library_workbook: str | Path,
) -> pd.DataFrame:
    """Read Table S4 and map author gRNA names to 10x guide identities."""

    name_mapping = _author_name_mapping(guide_library_workbook)
    rows: list[dict[str, object]] = []
    sheets = (("NANOG", "NANOG_GeneScores"), ("OCT4", "OCT4_GeneScores"))
    for target, sheet_name in sheets:
        headers, records = _sheet_records(dosage_workbook, sheet_name)
        if len(headers) < 5:
            raise ValueError(f"Unexpected Table S4 layout in {sheet_name!r}: {headers}")
        for record in records:
            raw_name = str(record.get(headers[0], "")).strip()
            if not raw_name:
                continue
            identity = name_mapping.get(raw_name)
            if identity is None:
                raise KeyError(f"Table S4 gRNA {raw_name!r} could not be mapped through the guide library")
            dose_pct = float(record[headers[1]])
            rows.append(
                {
                    "dose_target": target,
                    "author_grna": raw_name,
                    "guide_identity": identity,
                    "guide_family": guide_family(identity),
                    "is_control": guide_family(identity).startswith("control"),
                    "tf_expression_pct_control": dose_pct,
                    "tf_remaining": dose_pct / 100.0,
                    "knockdown_fraction": float(np.clip(1.0 - dose_pct / 100.0, 0.0, 1.0)),
                    "esc_score": float(record[headers[2]]),
                    "de_differentiation_score": float(record[headers[3]]),
                    "de_off_target_score": float(record[headers[4]]),
                }
            )
    table = pd.DataFrame(rows)
    if table.empty:
        raise ValueError(f"No dosage rows were read from {dosage_workbook}")
    duplicate_target_rows = table.duplicated(["dose_target", "guide_identity"])
    if duplicate_target_rows.any():
        duplicated = table.loc[duplicate_target_rows, ["dose_target", "guide_identity"]].to_dict("records")
        raise ValueError(f"Duplicate target/guide rows in Table S4: {duplicated}")
    return table.sort_values(["dose_target", "is_control", "tf_remaining", "guide_identity"]).reset_index(drop=True)


def validate_dosage_coverage(
    dosage: pd.DataFrame,
    guide_identities: Iterable[str],
    *,
    expected_total_guides: int = 48,
    expected_targeting_guides: int = 33,
) -> dict[str, object]:
    observed = {str(value) for value in guide_identities}
    dosage_guides = set(dosage["guide_identity"].astype(str))
    targeting = set(dosage.loc[~dosage["is_control"], "guide_identity"].astype(str))
    missing = sorted(observed - dosage_guides)
    extra = sorted(dosage_guides - observed)
    if len(observed) != expected_total_guides or len(targeting) != expected_targeting_guides or missing or extra:
        raise ValueError(
            "Dose coverage mismatch: "
            f"observed={len(observed)}, targeting={len(targeting)}, missing={missing}, extra={extra}"
        )
    return {
        "n_observed_guides": len(observed),
        "n_dosage_guides": len(dosage_guides),
        "n_targeting_guides": len(targeting),
        "n_target_guide_rows": int((~dosage["is_control"]).sum()),
        "n_control_target_rows": int(dosage["is_control"].sum()),
    }


def read_author_model_selection(dosage_workbook: str | Path) -> pd.DataFrame:
    headers, records = _sheet_records(dosage_workbook, "Gene_Score_Model_Selection")
    table = pd.DataFrame(records)
    if table.empty:
        raise ValueError("Gene_Score_Model_Selection is empty")
    return table.rename(columns={headers[0]: "name", headers[-1]: "model"})


def linear_curve(x: np.ndarray, intercept: float, slope: float) -> np.ndarray:
    return intercept + slope * x


def quadratic_curve(x: np.ndarray, intercept: float, linear: float, quadratic: float) -> np.ndarray:
    return intercept + linear * x + quadratic * x * x


def hill_curve(x: np.ndarray, ed50: float, hill: float) -> np.ndarray:
    x_safe = np.clip(np.asarray(x, dtype=float), 0.0, None)
    return 1.0 / (1.0 + np.power(x_safe / max(ed50, 1e-8), hill))


def _information_criteria(rss: float, n: int, k: int) -> tuple[float, float]:
    rss = max(float(rss), np.finfo(float).tiny)
    likelihood_k = k + 1  # Include the fitted residual variance, matching Table S4.
    aic = float(n * (math.log(2.0 * math.pi) + 1.0 + math.log(rss / n)) + 2 * likelihood_k)
    aicc = (
        float(aic + (2 * likelihood_k * (likelihood_k + 1)) / (n - likelihood_k - 1))
        if n > likelihood_k + 1
        else float("inf")
    )
    return aic, aicc


def _fit_parameters(model: str, x: np.ndarray, y: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    if model == "linear":
        params = np.polyfit(x, y, 1)[::-1]
        return params, linear_curve(x, *params)
    if model == "quadratic":
        params = np.polyfit(x, y, 2)[::-1]
        return params, quadratic_curve(x, *params)
    if model != "hill":
        raise ValueError(f"Unknown response model: {model}")
    positive_x = x[x > 0]
    min_ed50 = max(float(np.min(positive_x)) * 0.25, 1e-4) if positive_x.size else 1e-4
    max_ed50 = max(float(np.max(x)) * 2.0, min_ed50 * 10.0)

    def solve_grid(ed50_values: np.ndarray, coefficient_values: np.ndarray) -> tuple[float, np.ndarray]:
        x_grid = np.clip(x, 0.0, None)[None, None, :]
        ed50_grid = ed50_values[:, None, None]
        coefficient_grid = coefficient_values[None, :, None]
        basis = 1.0 / (1.0 + np.power(x_grid / ed50_grid, coefficient_grid))
        rss = np.square(y[None, None, :] - basis).sum(axis=2)
        best = np.unravel_index(int(np.argmin(rss)), rss.shape)
        params = np.asarray([ed50_values[best[0]], coefficient_values[best[1]]], dtype=float)
        return float(rss[best]), params

    best_rss = float("inf")
    best_params: np.ndarray | None = None
    ed50_grid = np.geomspace(min_ed50, max_ed50, 50)
    coefficient_grid = np.geomspace(0.1, 20.0, 50)
    best_rss, best_params = solve_grid(ed50_grid, coefficient_grid)
    if best_params is None:
        raise RuntimeError("Hill grid search failed")
    for _ in range(2):
        ed50_grid = np.geomspace(
            max(min_ed50, float(best_params[0]) / 1.8),
            min(max_ed50, float(best_params[0]) * 1.8),
            24,
        )
        coefficient_grid = np.geomspace(
            max(0.1, float(best_params[1]) / 1.8),
            min(20.0, float(best_params[1]) * 1.8),
            24,
        )
        rss, params = solve_grid(ed50_grid, coefficient_grid)
        if rss < best_rss:
            best_rss = rss
            best_params = params
    return best_params, hill_curve(x, *best_params)


def _loo_rmse(model: str, x: np.ndarray, y: np.ndarray) -> float:
    predictions: list[float] = []
    truths: list[float] = []
    for idx in range(len(x)):
        keep = np.arange(len(x)) != idx
        try:
            params, _ = _fit_parameters(model, x[keep], y[keep])
            function = {"linear": linear_curve, "quadratic": quadratic_curve, "hill": hill_curve}[model]
            predictions.append(float(function(np.asarray([x[idx]]), *params)[0]))
            truths.append(float(y[idx]))
        except (RuntimeError, ValueError, FloatingPointError, np.linalg.LinAlgError):
            return float("inf")
    return float(np.sqrt(np.mean(np.square(np.asarray(predictions) - np.asarray(truths)))))


def fit_response_models(
    x: Sequence[float],
    y: Sequence[float],
    *,
    include_loo: bool = True,
    models: Sequence[str] = ("linear", "quadratic", "hill"),
) -> list[dict[str, object]]:
    x_array = np.asarray(x, dtype=float)
    y_array = np.asarray(y, dtype=float)
    finite = np.isfinite(x_array) & np.isfinite(y_array)
    x_array = x_array[finite]
    y_array = y_array[finite]
    if len(x_array) < 6 or np.unique(x_array).size < 4:
        raise ValueError("Response fitting requires at least six finite rows and four distinct doses")
    order = np.argsort(x_array)
    x_array = x_array[order]
    y_array = y_array[order]
    rows: list[dict[str, object]] = []
    parameter_names = {
        "linear": ("intercept", "slope"),
        "quadratic": ("intercept", "linear", "quadratic"),
        "hill": ("ed50", "hill_coefficient"),
    }
    for model in models:
        try:
            params, fitted = _fit_parameters(model, x_array, y_array)
            rss = float(np.square(y_array - fitted).sum())
            aic, aicc = _information_criteria(rss, len(x_array), len(params))
            row: dict[str, object] = {
                "model": model,
                "n": len(x_array),
                "rss": rss,
                "aic": aic,
                "aicc": aicc,
                "loo_rmse": _loo_rmse(model, x_array, y_array) if include_loo else float("nan"),
                "fit_ok": True,
            }
            row.update({name: float(value) for name, value in zip(parameter_names[model], params)})
            if model == "quadratic":
                curvature = float(params[2])
                vertex = float(-params[1] / (2.0 * curvature)) if abs(curvature) > 1e-12 else float("nan")
                row["vertex"] = vertex
                row["inverted_u"] = bool(curvature < 0 and float(x_array.min()) < vertex < float(x_array.max()))
            elif model == "hill":
                row["threshold_in_observed_range"] = bool(
                    float(x_array.min()) < float(params[0]) < float(x_array.max())
                )
            rows.append(row)
        except (RuntimeError, ValueError, FloatingPointError, np.linalg.LinAlgError) as exc:
            rows.append(
                {
                    "model": model,
                    "n": len(x_array),
                    "rss": float("inf"),
                    "aic": float("inf"),
                    "aicc": float("inf"),
                    "loo_rmse": float("inf"),
                    "fit_ok": False,
                    "error": str(exc),
                }
            )
    finite_rows = [row for row in rows if np.isfinite(float(row["aic"]))]
    if not finite_rows:
        raise RuntimeError("All response models failed")
    best_aic = min(float(row["aic"]) for row in finite_rows)
    best_aicc = min(float(row["aicc"]) for row in finite_rows)
    for row in rows:
        row["delta_aic"] = float(row["aic"]) - best_aic
        row["delta_aicc"] = float(row["aicc"]) - best_aicc
        row["selected_aic"] = bool(float(row["aic"]) == best_aic)
        row["selected_aicc"] = bool(float(row["aicc"]) == best_aicc)
    return rows


def bootstrap_response_fit(
    x: Sequence[float],
    y: Sequence[float],
    *,
    n_bootstrap: int = 1000,
    seed: int = 7,
    models: Sequence[str] = ("linear", "quadratic", "hill"),
) -> dict[str, object]:
    x_array = np.asarray(x, dtype=float)
    y_array = np.asarray(y, dtype=float)
    finite = np.isfinite(x_array) & np.isfinite(y_array)
    x_array = x_array[finite]
    y_array = y_array[finite]
    rng = np.random.default_rng(seed)
    selected: list[str] = []
    vertices: list[float] = []
    ed50s: list[float] = []
    for _ in range(n_bootstrap):
        indices = rng.integers(0, len(x_array), size=len(x_array))
        if np.unique(x_array[indices]).size < 4:
            continue
        try:
            fits = fit_response_models(
                x_array[indices],
                y_array[indices],
                include_loo=False,
                models=models,
            )
        except (RuntimeError, ValueError):
            continue
        winner = next(row for row in fits if row["selected_aic"])
        selected.append(str(winner["model"]))
        quadratic = next((row for row in fits if row["model"] == "quadratic"), None)
        hill = next((row for row in fits if row["model"] == "hill"), None)
        if quadratic is not None and quadratic.get("inverted_u") and np.isfinite(float(quadratic.get("vertex", np.nan))):
            vertices.append(float(quadratic["vertex"]))
        if hill is not None and np.isfinite(float(hill.get("ed50", np.nan))):
            ed50s.append(float(hill["ed50"]))

    def interval(values: list[float]) -> list[float] | None:
        if not values:
            return None
        return [float(np.quantile(values, 0.025)), float(np.quantile(values, 0.975))]

    counts = {model: selected.count(model) for model in ("linear", "quadratic", "hill")}
    denominator = max(len(selected), 1)
    return {
        "requested": int(n_bootstrap),
        "successful": len(selected),
        "selection_frequency": {model: count / denominator for model, count in counts.items()},
        "quadratic_vertex_ci95": interval(vertices),
        "hill_ed50_ci95": interval(ed50s),
    }


def benjamini_hochberg(p_values: Sequence[float]) -> np.ndarray:
    values = np.asarray(p_values, dtype=float)
    result = np.full(values.shape, np.nan, dtype=float)
    finite_indices = np.flatnonzero(np.isfinite(values))
    if finite_indices.size == 0:
        return result
    ordered = finite_indices[np.argsort(values[finite_indices])]
    adjusted = values[ordered] * len(ordered) / np.arange(1, len(ordered) + 1)
    adjusted = np.minimum.accumulate(adjusted[::-1])[::-1]
    result[ordered] = np.clip(adjusted, 0.0, 1.0)
    return result


def linear_effect_summary(x: Sequence[float], y: Sequence[float]) -> dict[str, float]:
    x_array = np.asarray(x, dtype=float)
    y_array = np.asarray(y, dtype=float)
    finite = np.isfinite(x_array) & np.isfinite(y_array)
    result = stats.linregress(x_array[finite], y_array[finite])
    mean_expression = float(np.mean(y_array[finite]))
    scale = max(abs(mean_expression), 1e-6)
    dynamic_range = float(np.max(y_array[finite]) - np.min(y_array[finite]))
    return {
        "slope": float(result.slope),
        "mean_expression": mean_expression,
        "standardized_slope": float(result.slope / scale),
        "linear_pvalue": float(result.pvalue),
        "dynamic_range": dynamic_range,
        "relative_dynamic_range": float(dynamic_range / scale),
    }


def read_gene_sets(path: str | Path, available_genes: Iterable[str] | None = None) -> dict[str, list[str]]:
    raw = pd.read_excel(path, sheet_name=0, header=None)
    header_idx = next(
        idx
        for idx in range(raw.shape[0])
        if any("genes" in str(value).lower() or str(value).strip() == "Essentiality" for value in raw.iloc[idx])
    )
    available = None if available_genes is None else {str(value) for value in available_genes}
    aliases = {
        "pluripotency": "pluripotency",
        "differentiation": "differentiation",
        "de genes": "de",
        "ne genes": "ne",
        "essentiality": "essentiality",
    }
    gene_sets: dict[str, list[str]] = {}
    for col_idx, header in enumerate(raw.iloc[header_idx].tolist()):
        normalized = str(header or "").strip().replace("\n", " ").lower()
        key = next((value for prefix, value in aliases.items() if normalized.startswith(prefix)), None)
        if key is None:
            continue
        genes = []
        for value in raw.iloc[header_idx + 1 :, col_idx].dropna().tolist():
            gene = str(value).strip()
            if gene and (available is None or gene in available):
                genes.append(gene)
        gene_sets[key] = sorted(set(genes))
    return gene_sets


def build_rank_stratified_guide_folds(
    dosage: pd.DataFrame,
    *,
    n_folds: int = 5,
) -> dict[str, int]:
    if n_folds < 3:
        raise ValueError("At least three folds are required")
    targeting = dosage.loc[~dosage["is_control"]].copy()
    assignments: dict[str, int] = {}
    for target, group in targeting.groupby("dose_target", sort=True):
        ordered = group.sort_values(["knockdown_fraction", "guide_identity"])
        for rank, guide in enumerate(ordered["guide_identity"].astype(str)):
            assignments[guide] = rank % n_folds
        used = {assignments[str(guide)] for guide in ordered["guide_identity"]}
        if used != set(range(n_folds)):
            raise ValueError(f"Target {target!r} does not populate all {n_folds} folds")
    return assignments


def build_endpoint_interpolation_partition(
    dosage: pd.DataFrame,
    *,
    endpoint_fraction: float = 0.30,
) -> dict[str, set[str]]:
    """Hold out the middle of each target's dose range at guide level.

    The outer dose bands form the training domain.  The guide immediately
    inside each extreme is reserved for validation, so model selection never
    observes a middle-dose guide.  All remaining middle guides are test-only.
    """

    if not 0.0 < endpoint_fraction < 0.5:
        raise ValueError("endpoint_fraction must be between 0 and 0.5")
    targeting = dosage.loc[~dosage["is_control"]].copy()
    partitions = {"train": set(), "val": set(), "test": set()}
    for target, group in targeting.groupby("dose_target", sort=True):
        ordered = group.sort_values(["knockdown_fraction", "guide_identity"])
        guides = ordered["guide_identity"].astype(str).tolist()
        n_endpoint = max(2, int(np.ceil(len(guides) * endpoint_fraction)))
        if 2 * n_endpoint >= len(guides):
            raise ValueError(
                f"Target {target!r} has too few guides for endpoint interpolation: {len(guides)}"
            )
        low = guides[:n_endpoint]
        high = guides[-n_endpoint:]
        partitions["val"].update((low[-1], high[0]))
        partitions["train"].update(low[:-1])
        partitions["train"].update(high[1:])
        partitions["test"].update(guides[n_endpoint:-n_endpoint])
    validate_fold_partition(partitions["train"], partitions["val"], partitions["test"])
    return partitions


def validate_fold_partition(
    train_guides: Iterable[str],
    val_guides: Iterable[str],
    test_guides: Iterable[str],
) -> None:
    train = set(train_guides)
    val = set(val_guides)
    test = set(test_guides)
    if train & val or train & test or val & test:
        raise ValueError("Guide leakage detected across train/validation/test partitions")
