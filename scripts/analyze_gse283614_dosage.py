from __future__ import annotations

import argparse
import csv
from dataclasses import replace
import json
from pathlib import Path
import sys

import numpy as np
from scipy import sparse
import torch

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "src"))

from perturbdiff.gse283614 import guide_family, guide_target_gene
from perturbdiff.data import canonical_guide_keys, load_guide_activity_lookup
from perturbdiff.model import PerturbDiffusionModel


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Biological dose-response analysis for prepared GSE283614 h5ad files.")
    parser.add_argument("--adata-path", type=Path, required=True)
    parser.add_argument("--checkpoint", type=Path, default=None)
    parser.add_argument("--checkpoint-name", default="perturbdiff")
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--output-prefix", default="gse283614_dosage")
    parser.add_argument("--perturbation-col", default="perturbation_label")
    parser.add_argument("--guide-col", default="guide_identity")
    parser.add_argument("--state-col", default="cell_state")
    parser.add_argument("--min-cells-per-guide", type=int, default=10)
    parser.add_argument("--max-guides", type=int, default=0)
    parser.add_argument("--guide-offset", type=int, default=0)
    parser.add_argument("--gene-set-workbook", type=Path, default=None)
    parser.add_argument("--reference-dosage-workbook", type=Path, default=None)
    parser.add_argument("--activity-table", type=Path, default=None)
    parser.add_argument("--activity-col", default="relative_activity")
    parser.add_argument("--max-control-cells", type=int, default=128)
    parser.add_argument("--ddim-steps", type=int, default=25)
    parser.add_argument("--top-genes-per-guide", type=int, default=25)
    parser.add_argument("--seed", type=int, default=7)
    parser.add_argument("--device", default="cuda" if torch.cuda.is_available() else "cpu")
    return parser.parse_args()


def to_dense(x) -> np.ndarray:
    if sparse.issparse(x):
        return x.toarray()
    return np.asarray(x)


def vector_mean(x) -> np.ndarray:
    if sparse.issparse(x):
        return np.asarray(x.mean(axis=0)).ravel()
    return np.asarray(x).mean(axis=0)


def read_gene_sets(path: Path | None, gene_names: list[str]) -> dict[str, list[str]]:
    defaults = {
        "pluripotency": ["POU5F1", "NANOG", "SOX2"],
        "differentiation": ["SOX17", "CER1", "FOXA2", "GATA6", "EOMES"],
    }
    if path is None or not path.exists():
        return {
            name: [gene for gene in genes if gene in gene_names]
            for name, genes in defaults.items()
        }
    try:
        import openpyxl
    except ImportError as exc:
        raise ImportError("Reading gene-set workbooks requires openpyxl") from exc

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = workbook.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header_idx = None
    for idx, row in enumerate(rows):
        values = [str(value).strip() for value in row if value is not None]
        if any("genes" in value.lower() for value in values) or "Essentiality" in values:
            header_idx = idx
            break
    if header_idx is None:
        return defaults
    headers = [str(value).strip().replace("\n", " ") if value is not None else "" for value in rows[header_idx]]
    gene_sets: dict[str, list[str]] = {}
    for col_idx, header in enumerate(headers):
        if not header:
            continue
        genes: list[str] = []
        for row in rows[header_idx + 1 :]:
            if col_idx >= len(row):
                continue
            value = row[col_idx]
            if value is None:
                continue
            gene = str(value).strip()
            if gene and gene in gene_names:
                genes.append(gene)
        if genes:
            key = header.lower().replace(" ", "_")
            gene_sets[key] = sorted(set(genes))
    return gene_sets or defaults


def load_reference_dosage_tables(path: Path | None) -> dict[str, object]:
    if path is None or not path.exists():
        return {}
    try:
        import openpyxl
    except ImportError as exc:
        raise ImportError("Reading dosage workbooks requires openpyxl") from exc

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    selected = {"NANOG_GeneScores", "OCT4_GeneScores", "Gene_Score_Model_Selection"}
    out: dict[str, object] = {}
    for ws in workbook.worksheets:
        if ws.title not in selected:
            continue
        rows = list(ws.iter_rows(values_only=True))
        header_idx = None
        for idx, row in enumerate(rows):
            first = "" if row[0] is None else str(row[0]).strip()
            if first in {"gRNA", "Name"}:
                header_idx = idx
                break
        if header_idx is None:
            continue
        headers = ["" if value is None else str(value).strip() for value in rows[header_idx]]
        records: list[dict[str, object]] = []
        for row in rows[header_idx + 1 :]:
            if not any(value is not None for value in row):
                continue
            records.append(
                {
                    headers[idx]: row[idx]
                    for idx in range(min(len(headers), len(row)))
                    if headers[idx]
                }
            )
        out[ws.title] = {"n_rows": len(records), "columns": headers, "records": records}
    return out


def compatible_config_from_state(state: dict) -> object:
    cfg = state["model_config"]
    state_dict = state["model_state_dict"]
    has_activity_condition = any(key.startswith("activity_condition.") for key in state_dict)
    if getattr(cfg.guide_encoder, "enabled", False):
        cfg = replace(
            cfg,
            guide_encoder=replace(
                cfg.guide_encoder,
                explicit_condition=has_activity_condition,
                activity_scales_delta=not has_activity_condition,
            ),
        )
    return cfg


def load_model(path: Path, device: torch.device) -> tuple[PerturbDiffusionModel, dict[str, int]]:
    state = torch.load(path, map_location=device, weights_only=False)
    model = PerturbDiffusionModel(compatible_config_from_state(state)).to(device)
    model.load_state_dict(state["model_state_dict"], strict=False)
    model.eval()
    return model, {str(key): int(value) for key, value in state["perturbation_mapping"].items()}


def gene_indices(gene_names: list[str], genes: list[str]) -> list[int]:
    index = {gene: idx for idx, gene in enumerate(gene_names)}
    return [index[gene] for gene in genes if gene in index]


def pct_control(mean_vec: np.ndarray, control_mean: np.ndarray, gene_idx: int | None) -> float:
    if gene_idx is None:
        return float("nan")
    denom = float(control_mean[gene_idx])
    if abs(denom) < 1e-8:
        return float("nan")
    return float(100.0 * mean_vec[gene_idx] / denom)


def module_score(mean_vec: np.ndarray, indices: list[int]) -> float:
    if not indices:
        return float("nan")
    return float(np.mean(mean_vec[indices]))


def add_scores(
    row: dict[str, object],
    mean_vec: np.ndarray,
    control_mean: np.ndarray,
    *,
    gene_names: list[str],
    gene_sets: dict[str, list[str]],
) -> None:
    index = {gene: idx for idx, gene in enumerate(gene_names)}
    row["POU5F1_pct_control"] = pct_control(mean_vec, control_mean, index.get("POU5F1"))
    row["NANOG_pct_control"] = pct_control(mean_vec, control_mean, index.get("NANOG"))
    for name, genes in gene_sets.items():
        indices = gene_indices(gene_names, genes)
        score = module_score(mean_vec, indices)
        control_score = module_score(control_mean, indices)
        row[f"{name}_n_genes"] = len(indices)
        row[f"{name}_score"] = score
        row[f"{name}_delta_control"] = score - control_score if np.isfinite(score) and np.isfinite(control_score) else float("nan")


def ranking_columns(gene_sets: dict[str, list[str]]) -> tuple[str | None, str | None]:
    plur = next((key for key in gene_sets if "pluripotency" in key), None)
    diff = next((key for key in gene_sets if "differentiation" in key), None)
    return plur, diff


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    fields = sorted({key for row in rows for key in row})
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


@torch.no_grad()
def predict_mean(
    model: PerturbDiffusionModel,
    mapping: dict[str, int],
    control_cells: np.ndarray,
    *,
    guide_label: str,
    target_label: str,
    activity_value: float | None,
    steps: int,
    device: torch.device,
) -> tuple[np.ndarray | None, float | None]:
    target_id = mapping.get(str(guide_label), mapping.get(str(target_label)))
    if target_id is None:
        return None, None
    control_tensor = torch.from_numpy(control_cells.astype(np.float32, copy=False)).to(device)
    target_ids = torch.full((control_tensor.shape[0],), int(target_id), dtype=torch.long, device=device)
    guide_features = None
    if getattr(model, "guide_doser", None) is not None:
        if activity_value is None:
            raise ValueError(f"Activity-enabled checkpoint requires a mapped activity for {guide_label!r}")
        guide_features = torch.full(
            (control_tensor.shape[0], 1),
            float(activity_value),
            dtype=control_tensor.dtype,
            device=device,
        )
    pred = model.sample_counterfactual(
        control_tensor,
        target_ids,
        guide_features=guide_features,
        steps=steps,
    )
    norm = float(model.perturbation.embedding.weight[int(target_id)].detach().norm().cpu())
    return pred.detach().cpu().numpy().mean(axis=0), norm


def main() -> None:
    try:
        import anndata as ad
    except ImportError as exc:
        raise ImportError("analyze_gse283614_dosage requires anndata") from exc

    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng(args.seed)
    torch.manual_seed(args.seed)

    adata = ad.read_h5ad(args.adata_path)
    for col in (args.perturbation_col, args.guide_col, args.state_col):
        if col not in adata.obs:
            raise KeyError(f"Column {col!r} not found in adata.obs")
    gene_names = list(adata.var_names.astype(str))
    gene_sets = read_gene_sets(args.gene_set_workbook, gene_names)
    plur_key, diff_key = ranking_columns(gene_sets)

    model = None
    mapping: dict[str, int] = {}
    device = torch.device(args.device)
    if args.checkpoint is not None:
        model, mapping = load_model(args.checkpoint, device)
    activity_lookup: dict[str, float] = {}
    if model is not None and getattr(model, "guide_doser", None) is not None:
        if args.activity_table is None:
            raise ValueError("Activity-enabled checkpoints require --activity-table and --activity-col")
        activity_lookup = load_guide_activity_lookup(args.activity_table, activity_col=args.activity_col)

    labels = adata.obs[args.perturbation_col].astype(str).to_numpy()
    guides = adata.obs[args.guide_col].astype(str).to_numpy()
    states = adata.obs[args.state_col].astype(str).to_numpy()
    score_rows: list[dict[str, object]] = []
    top_gene_rows: list[dict[str, object]] = []

    for state in sorted(np.unique(states).tolist()):
        state_mask = states == state
        control_mask = state_mask & (labels == "control")
        if int(control_mask.sum()) == 0:
            continue
        control_mean = vector_mean(adata.X[control_mask])
        control_cells = to_dense(adata.X[control_mask]).astype(np.float32, copy=False)
        if args.max_control_cells > 0 and control_cells.shape[0] > args.max_control_cells:
            control_cells = control_cells[rng.choice(control_cells.shape[0], args.max_control_cells, replace=False)]

        state_guides = []
        for guide in sorted(set(guides[state_mask].tolist())):
            mask = state_mask & (guides == guide) & (labels != "control")
            if int(mask.sum()) >= args.min_cells_per_guide:
                state_guides.append(str(guide))
        state_guides = state_guides[args.guide_offset :]
        if args.max_guides > 0:
            state_guides = state_guides[: args.max_guides]

        for guide in state_guides:
            mask = state_mask & (guides == guide)
            true_cells = adata.X[mask]
            observed_mean = vector_mean(true_cells)
            family = str(adata.obs.loc[mask, "guide_family"].iloc[0]) if "guide_family" in adata.obs else guide_family(guide)
            target = str(adata.obs.loc[mask, "target_gene"].iloc[0]) if "target_gene" in adata.obs else guide_target_gene(guide)
            base = {
                "state": state,
                "guide_identity": guide,
                "perturbation_label": str(adata.obs.loc[mask, args.perturbation_col].iloc[0]),
                "guide_family": family,
                "target_gene": target,
                "n_cells": int(mask.sum()),
            }
            observed_row = {**base, "source": "observed"}
            add_scores(observed_row, observed_mean, control_mean, gene_names=gene_names, gene_sets=gene_sets)
            if plur_key is not None:
                observed_row["pluripotency_drop"] = -float(observed_row.get(f"{plur_key}_delta_control", float("nan")))
            if diff_key is not None:
                observed_row["differentiation_gain"] = float(observed_row.get(f"{diff_key}_delta_control", float("nan")))
            score_rows.append(observed_row)

            true_delta = observed_mean - control_mean
            top_indices = np.argsort(np.abs(true_delta))[::-1][: args.top_genes_per_guide]
            for rank, gene_idx in enumerate(top_indices, start=1):
                top_gene_rows.append(
                    {
                        **base,
                        "source": "observed",
                        "rank": rank,
                        "gene": gene_names[int(gene_idx)],
                        "delta_control": float(true_delta[int(gene_idx)]),
                    }
                )

            if model is None:
                continue
            activity_value = None
            if getattr(model, "guide_doser", None) is not None:
                activity_value = next(
                    (activity_lookup[key] for key in canonical_guide_keys(guide) if key in activity_lookup),
                    None,
                )
            pred_mean, embedding_norm = predict_mean(
                model,
                mapping,
                control_cells,
                guide_label=guide,
                target_label=family,
                activity_value=activity_value,
                steps=args.ddim_steps,
                device=device,
            )
            if pred_mean is None:
                continue
            pred_row = {**base, "source": args.checkpoint_name, "embedding_l2_norm": embedding_norm}
            add_scores(pred_row, pred_mean, control_mean, gene_names=gene_names, gene_sets=gene_sets)
            if plur_key is not None:
                pred_row["pluripotency_drop"] = -float(pred_row.get(f"{plur_key}_delta_control", float("nan")))
            if diff_key is not None:
                pred_row["differentiation_gain"] = float(pred_row.get(f"{diff_key}_delta_control", float("nan")))
            score_rows.append(pred_row)

            pred_delta = pred_mean - control_mean
            top_indices = np.argsort(np.abs(pred_delta))[::-1][: args.top_genes_per_guide]
            for rank, gene_idx in enumerate(top_indices, start=1):
                top_gene_rows.append(
                    {
                        **base,
                        "source": args.checkpoint_name,
                        "rank": rank,
                        "gene": gene_names[int(gene_idx)],
                        "delta_control": float(pred_delta[int(gene_idx)]),
                    }
                )

    guide_scores_path = args.output_dir / f"{args.output_prefix}_guide_scores.csv"
    top_genes_path = args.output_dir / f"{args.output_prefix}_top_genes.csv"
    reference_path = args.output_dir / f"{args.output_prefix}_reference_tables.json"
    summary_path = args.output_dir / f"{args.output_prefix}_summary.json"
    write_csv(guide_scores_path, score_rows)
    write_csv(top_genes_path, top_gene_rows)
    reference = load_reference_dosage_tables(args.reference_dosage_workbook)
    reference_path.write_text(json.dumps(reference, indent=2, default=str), encoding="utf-8")

    summary = {
        "adata_path": str(args.adata_path),
        "checkpoint": None if args.checkpoint is None else str(args.checkpoint),
        "n_rows": len(score_rows),
        "n_top_gene_rows": len(top_gene_rows),
        "gene_sets": {name: len(genes) for name, genes in gene_sets.items()},
        "outputs": {
            "guide_scores": str(guide_scores_path),
            "top_genes": str(top_genes_path),
            "reference_tables": str(reference_path),
        },
    }
    summary_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
